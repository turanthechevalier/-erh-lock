"""Pandoc'un varsayılan referans belgesinden kitap biçimli bir referans DOCX üretir.

Sayfa: 16 x 24 cm (Türk akademik yayıncılığında yaygın kitap ebadı)
Yazı: Liberation Serif (Times New Roman ile ölçü uyumlu), metin 11 pt, dipnot 9 pt.
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

KAYNAK, HEDEF = sys.argv[1], sys.argv[2]
YAZI = "Liberation Serif"

belge = Document(KAYNAK)

for bolum in belge.sections:
    bolum.page_width = Cm(16)
    bolum.page_height = Cm(24)
    bolum.top_margin = Cm(2.2)
    bolum.bottom_margin = Cm(2.2)
    bolum.left_margin = Cm(2.2)
    bolum.right_margin = Cm(2.0)
    bolum.header_distance = Cm(1.0)
    bolum.footer_distance = Cm(1.0)


def stil_bul(ad):
    kimlik = ad.replace(" ", "")
    for s in belge.styles:
        if s.style_id == kimlik or (s.name or "").lower() == ad.lower():
            return s
    return None


def bicimle(ad, boyut, kalin=None, italik=None, hiza=None, once=None, sonra=None,
            satir=None, girinti=None):
    stil = stil_bul(ad)
    if stil is None:
        print("uyarı: stil bulunamadı:", ad)
        return
    stil.font.name = YAZI
    stil.font.size = Pt(boyut)
    if kalin is not None:
        stil.font.bold = kalin
    if italik is not None:
        stil.font.italic = italik
    pf = stil.paragraph_format
    if hiza is not None:
        pf.alignment = hiza
    if once is not None:
        pf.space_before = Pt(once)
    if sonra is not None:
        pf.space_after = Pt(sonra)
    if satir is not None:
        pf.line_spacing = satir
    if girinti is not None:
        pf.first_line_indent = Cm(girinti)
    # Doğu Asya / karmaşık yazı fontlarını da eşitle
    rpr = stil.element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        for anahtar in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + anahtar, YAZI)
        for anahtar in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            rfonts.attrib.pop("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + anahtar, None)


J = WD_ALIGN_PARAGRAPH.JUSTIFY
C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT

bicimle("Normal", 11, hiza=J, sonra=4, satir=1.15)
bicimle("Body Text", 11, hiza=J, once=0, sonra=4, satir=1.15, girinti=0.7)
bicimle("First Paragraph", 11, hiza=J, once=0, sonra=4, satir=1.15, girinti=0.7)
bicimle("Compact", 9, hiza=L, once=0, sonra=1, satir=1.0, girinti=0)
bicimle("Block Text", 10, hiza=J, once=4, sonra=6, satir=1.1)
bicimle("Footnote Text", 9, hiza=J, once=0, sonra=1, satir=1.0)
bicimle("Title", 20, kalin=True, hiza=C, once=60, sonra=18)
bicimle("Subtitle", 13, italik=True, hiza=C, once=6, sonra=30)
bicimle("Author", 12, hiza=C, once=24, sonra=6)
bicimle("Date", 11, hiza=C)
bicimle("Heading 1", 15, kalin=True, hiza=C, once=36, sonra=18)
bicimle("Heading 2", 12.5, kalin=True, hiza=L, once=18, sonra=8)
bicimle("Heading 3", 11.5, kalin=True, hiza=L, once=12, sonra=6)
bicimle("Heading 4", 11, kalin=True, italik=False, hiza=L, once=10, sonra=4)
bicimle("Heading 5", 11, kalin=False, italik=True, hiza=L, once=8, sonra=4)
bicimle("Table", 9.5, hiza=L)
bicimle("Bibliography", 10, hiza=J)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
for ad in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Title",
           "Subtitle", "TOC Heading"):
    stil = stil_bul(ad)
    if stil is None:
        continue
    rpr = stil.element.get_or_add_rPr()
    for renk in rpr.findall(W + "color"):
        rpr.remove(renk)
    stil.font.color.rgb = RGBColor(0, 0, 0)

bicimle("TOC Heading", 14, kalin=True, hiza=C, once=0, sonra=18)
for ad in ("Heading 1", "TOC Heading"):
    stil = stil_bul(ad)
    if stil is not None:
        stil.paragraph_format.page_break_before = True
for ad in ("Heading 2", "Heading 3", "Heading 4", "Heading 5"):
    stil = stil_bul(ad)
    if stil is not None:
        stil.paragraph_format.keep_with_next = True

# Alt bilgiye ortalanmış sayfa numarası
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

for bolum in belge.sections:
    altbilgi = bolum.footer
    altbilgi.is_linked_to_previous = False
    p = altbilgi.paragraphs[0] if altbilgi.paragraphs else altbilgi.add_paragraph()
    p.alignment = C
    calisma = p.add_run()
    calisma.font.name = YAZI
    calisma.font.size = Pt(9)
    for tur, metin in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if tur:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tur)
            calisma._r.append(fld)
        elif metin == " PAGE ":
            komut = OxmlElement("w:instrText")
            komut.set(qn("xml:space"), "preserve")
            komut.text = metin
            calisma._r.append(komut)
        else:
            t = OxmlElement("w:t")
            t.text = metin
            calisma._r.append(t)

belge.save(HEDEF)
print("referans belge yazıldı:", HEDEF)
